from pptx import Presentation
from pptx.util import Inches, Pt

def create_presentation():
    prs = Presentation()

    # Slide 1: Title
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]

    title.text = "Aegis.AI"
    subtitle.text = "Decentralized Threat Intelligence Against Scam Syndicates\n\nTeam:\nHARSHIT | SHOURYA SINGH | HARSH VARDHAN SINGH | SHIVAM RAJ"

    def add_bullet_slide(prs, title_text, bullet_points):
        bullet_slide_layout = prs.slide_layouts[1]
        slide = prs.slides.add_slide(bullet_slide_layout)
        shapes = slide.shapes
        title_shape = shapes.title
        body_shape = shapes.placeholders[1]
        
        title_shape.text = title_text
        
        tf = body_shape.text_frame
        if bullet_points:
            tf.text = bullet_points[0]
            for point in bullet_points[1:]:
                p = tf.add_paragraph()
                p.text = point
                p.level = 0

    # Slide 2: The Problem
    add_bullet_slide(prs, "The Problem", [
        "Scam syndicates use encrypted channels (Telegram/WhatsApp) to distribute counterfeit trading apps.",
        "Apps perfectly clone the UI of platforms like Zerodha and Groww.",
        "AI-generated deepfakes of financial leaders manipulate public sentiment.",
        "Centralized databases are slow, censorable, and vulnerable to poisoning attacks."
    ])

    # Slide 3: Our Solution: Aegis.AI
    add_bullet_slide(prs, "Our Solution: Aegis.AI", [
        "A decentralized, community-driven threat intelligence platform.",
        "Ingestion Layer: Scrapes Telegram for illicit APKs and synthetic media.",
        "Behavioral Engine: Static analysis of AndroidManifest.xml (flags SYSTEM_ALERT_WINDOW).",
        "Computer Vision: Siamese Neural Networks detect cross-platform UI clones.",
        "Deepfake Engine: Temporal optical flow (RAFT) & Face X-Ray boundary detection."
    ])

    # Slide 4: Token Curated Registry (TCR)
    add_bullet_slide(prs, "Decentralized Token Curated Registry (TCR)", [
        "Built on Ethereum (Solidity Smart Contracts).",
        "Cryptoeconomic Security: Researchers lock an Integrity Bond to submit an IoC.",
        "Community Consensus: Fraudulent submissions are challenged; a vote determines the outcome.",
        "Slashing Mechanism: Malicious actors lose their stake, rewarding honest verifiers.",
        "No Single Point of Failure: Trustless, immutable, and community-governed."
    ])

    # Slide 5: Regulatory Compliance (India DPDP Act)
    add_bullet_slide(prs, "Regulatory Compliance", [
        "Strict adherence to the India Digital Personal Data Protection (DPDP) Act, 2023.",
        "Publicly Available Data: Only parses open, public Telegram domains.",
        "Data Minimization: Zero PII storage. Artifacts are immediately hashed (SHA-256).",
        "Transparent Redressal: Blockchain provides an immutable 'Right to Correction'."
    ])

    # Slide 6: Conclusion
    add_bullet_slide(prs, "Conclusion", [
        "Aegis.AI makes it financially ruinous for scammers to operate.",
        "Scalable microservice architecture deployed in 48 hours.",
        "",
        "Thank You. Open for Q&A."
    ])

    prs.save("Aegis_AI_Pitch_Deck.pptx")

if __name__ == "__main__":
    create_presentation()
